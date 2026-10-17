"""Bounded local-file intake and conservative, part-preserving Office exports.

Source metadata is private repository data, never trusted from the browser. Office
patches change values/text only; structural changes require an explicit clean
conversion instead of silently discarding unsupported document features.
"""
from __future__ import annotations

import copy
import csv
import io
import posixpath
import re
import zipfile
from collections import Counter
from dataclasses import dataclass
from pathlib import PurePosixPath

from lxml import etree as ET

from project_docs import markdown_import
from project_docs.doc_schema import column_ids, new_id, node_text, validate_doc

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_CELLS = 20_000
MAX_ROWS = 10_000
MAX_COLUMNS = 100
MIME = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "csv": "text/csv; charset=utf-8", "md": "text/markdown; charset=utf-8"}
S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
XML = "http://www.w3.org/XML/1998/namespace"


class HandoffError(ValueError):
    pass


@dataclass
class ImportedFile:
    doc: dict
    surface_type: str
    filename: str
    mime: str
    metadata: dict


def _xml(raw):
    if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
        raise HandoffError("Office XML entities are not supported")
    return ET.fromstring(raw, parser=ET.XMLParser(resolve_entities=False, no_network=True))


def _archive(raw):
    try:
        archive = zipfile.ZipFile(io.BytesIO(raw))
        items = archive.infolist()
        if len(items) > 2000 or len({i.filename for i in items}) != len(items):
            raise HandoffError("Office archive has too many or duplicate entries")
        if sum(i.file_size for i in items) > MAX_EXPANDED_BYTES:
            raise HandoffError("Office archive expands beyond 64 MB")
        for item in items:
            if item.flag_bits & 1 or item.file_size > 16 * 1024 * 1024:
                raise HandoffError("Encrypted or oversized Office part")
            if item.file_size > 1_000_000 and item.file_size / max(item.compress_size, 1) > 200:
                raise HandoffError("Office archive compression ratio is unsafe")
            if ".." in PurePosixPath(item.filename).parts or item.filename.startswith(("/", "\\")):
                raise HandoffError("Invalid Office part path")
            if item.filename.endswith((".xml", ".rels")):
                part = archive.read(item.filename)
                if b"<!DOCTYPE" in part.upper() or b"<!ENTITY" in part.upper():
                    raise HandoffError("Office XML entities are not supported")
        return archive
    except (zipfile.BadZipFile, OSError, RuntimeError) as exc:
        raise HandoffError("Malformed Office archive") from exc


def _paragraph(text="", **attrs):
    node = {"type": "paragraph", "attrs": {"bid": new_id("b"), **attrs}}
    if text:
        node["content"] = [{"type": "text", "text": str(text)}]
    return node


def _table(rows, name=None):
    if len(rows) > MAX_ROWS or any(len(row) > MAX_COLUMNS for row in rows):
        raise HandoffError("Sheet exceeds 10,000 rows or 100 columns")
    if sum(len(row) for row in rows) > MAX_CELLS:
        raise HandoffError("Sheet exceeds 20,000 cells")
    attrs = {"bid": new_id("b")}
    if name is not None:
        attrs["worksheet_name"] = name
        attrs["sheet_name"] = name
    return {"type": "table", "attrs": attrs, "content": [
        {"type": "tableRow", "attrs": {"bid": new_id("b")}, "content": [
            {"type": "tableHeader" if i == 0 else "tableCell", "content": [_paragraph(value)]}
            for value in row]} for i, row in enumerate(rows)]}


def _tables(doc):
    return [n for n in doc.get("content", []) if n["type"] == "table"]


def _grid(table):
    return [[node_text(cell) for cell in row.get("content", [])]
            for row in table.get("content", [])]


def _xlsx(raw, *, header_row=1, worksheet=None):
    import openpyxl
    with _archive(raw) as archive:
        workbook = _xml(archive.read("xl/workbook.xml"))
        rels = _xml(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {e.get("Id"): e.get("Target") for e in rels if e.get("TargetMode") != "External"}
        worksheets = []
        for sheet in workbook.findall(f"{{{S}}}sheets/{{{S}}}sheet"):
            target = targets[sheet.get(f"{{{R}}}id")]
            path = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            worksheets.append((sheet.get("name"), path))
        # Preflight XML dimensions before openpyxl can iterate a forged huge range.
        for _, path in worksheets:
            root = _xml(archive.read(path))
            cells = root.findall(f".//{{{S}}}sheetData/{{{S}}}row/{{{S}}}c")
            if len(cells) > MAX_CELLS:
                raise HandoffError("Workbook exceeds 20,000 cells")
            for cell in cells:
                coordinate = cell.get("r", "A1")
                row, col = openpyxl.utils.cell.coordinate_to_tuple(coordinate)
                if row > MAX_ROWS or col > MAX_COLUMNS:
                    raise HandoffError("Workbook exceeds 10,000 rows or 100 columns")
        book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=True)
        content, mapping, total = [], [], 0
        try:
            for name, path in worksheets:
                sheet = book[name]
                if (sheet.max_row or 1) > MAX_ROWS or (sheet.max_column or 1) > MAX_COLUMNS:
                    raise HandoffError("Workbook dimensions exceed supported limits")
                total += (sheet.max_row or 1) * (sheet.max_column or 1)
                if total > MAX_CELLS:
                    raise HandoffError("Workbook exceeds 20,000 cells")
                rows = [["" if v is None else (v.isoformat() if hasattr(v, "isoformat") else str(v))
                         for v in row] for row in sheet.iter_rows(values_only=True)]
                start = header_row if worksheet in (None, name) else 1
                if not 1 <= start <= max(1, len(rows)):
                    raise HandoffError(f"{name}: header_row is outside the worksheet")
                table = _table(rows[start - 1:] or [[""]], name)
                table["attrs"].update({"source_header_row": start, "source_data_start_row": start + 1,
                                       "source_start_column": 1, "mapping_revision": 1})
                content.append(table)
                mapping.append({"name": name, "path": path, "header_row": start, "start_row": start,
                                "start_column": 1, "data_start_row": start + 1, "row_count": len(rows) - start + 1,
                                "column_count": sheet.max_column or 1})
        finally:
            book.close()
    return {"type": "doc", "content": content}, {"worksheets": mapping}


def _docx(raw):
    with _archive(raw) as archive:
        root = _xml(archive.read("word/document.xml"))
        body = root.find(f"{{{W}}}body")
        if body is None:
            raise HandoffError("DOCX has no document body")
        content, indices = [], []
        for index, child in enumerate(body):
            if child.tag == f"{{{W}}}tbl":
                grid = [["".join(cell.xpath(".//w:t/text()", namespaces={"w": W})) for cell in row.findall(f"{{{W}}}tc")]
                        for row in child.findall(f"{{{W}}}tr")]
                content.append(_table(grid))
                indices.append(index)
                continue
            if child.tag != f"{{{W}}}p":
                continue
            text = "".join(child.xpath(".//w:t/text()", namespaces={"w": W}))
            content.append(_paragraph(text))
            indices.append(index)
            if len(content) > MAX_ROWS:
                raise HandoffError("DOCX exceeds 10,000 paragraphs")
    return {"type": "doc", "content": content}, {"paragraph_indices": indices}


def import_file(filename: str, data: bytes, *, slug: str, header_row=1, worksheet=None, delivery_mode="original_patch") -> ImportedFile:
    filename = str(filename).replace("\\", "/").rsplit("/", 1)[-1].strip()
    if not filename or len(filename) > 200 or any(ord(c) < 32 for c in filename):
        raise HandoffError("Invalid filename")
    fmt = filename.rsplit(".", 1)[-1].lower()
    if fmt not in MIME:
        raise HandoffError("Supported formats: .xlsx, .docx, .csv, .md")
    if not data or len(data) > MAX_FILE_BYTES:
        raise HandoffError("File must contain between 1 byte and 8 MB")
    try:
        header_row = int(header_row)
        if header_row < 1 or delivery_mode not in ("original_patch", "new_workbook"):
            raise HandoffError("Invalid header row or delivery mode")
        if fmt == "xlsx":
            doc, metadata = _xlsx(data, header_row=header_row, worksheet=worksheet)
        elif fmt == "docx":
            doc, metadata = _docx(data)
        else:
            text = data.decode("utf-8-sig")
            if "\x00" in text:
                raise HandoffError("Text files must not contain NUL bytes")
            if fmt == "csv":
                reader = csv.reader(io.StringIO(text, newline=""), strict=True)
                rows = []
                total = 0
                for row in reader:
                    total += len(row)
                    if len(rows) >= MAX_ROWS or len(row) > MAX_COLUMNS or total > MAX_CELLS:
                        raise HandoffError("CSV exceeds supported row/column/cell limits")
                    rows.append(row)
                if header_row > max(1, len(rows)):
                    raise HandoffError("Header row is outside the CSV")
                doc, metadata = {"type": "doc", "content": [_table(rows[header_row - 1:] or [[""]])]}, {"header_row": header_row}
            else:
                doc, _ = markdown_import.markdown_to_doc(text, tab_keys=set(), image_src=lambda _: None,
                                                        fill_headers=())
                metadata = {}
        doc, _ = validate_doc(doc, slug=slug)
        metadata.update({"format": fmt, "baseline": copy.deepcopy(doc), "delivery_mode": delivery_mode, "mapping_revision": 1})
        return ImportedFile(doc, "sheet" if fmt in ("xlsx", "csv") else "doc", filename, MIME[fmt], metadata)
    except HandoffError:
        raise
    except Exception as exc:
        # Libraries produce many exception types for malformed XML and containers.
        raise HandoffError(f"Cannot read {fmt.upper()} file: malformed or unsupported content") from exc


def _identity(node):
    if node["type"] == "table":
        return node.get("attrs", {}).get("worksheet_name") or "table"
    return node_text(node)


def preview_file(filename, data, *, slug):
    imported = import_file(filename, data, slug=slug)
    return {"filename": imported.filename, "format": imported.metadata["format"],
            "worksheets": [{"name": table.get("attrs", {}).get("worksheet_name", ""),
                            "row_count": len(table.get("content", [])),
                            "column_count": max((len(r.get("content", [])) for r in table.get("content", [])), default=0),
                            "rows": _grid(table)[:20]} for table in _tables(imported.doc)],
            "warnings": ["僅整理選定表頭以後的資料。合併格、特殊公式與原格式並非完整可編輯；更動結構需另存新工作簿。"] if imported.metadata["format"] == "xlsx" else []}


def compatibility(doc, source):
    metadata = source.get("metadata") or {}
    fmt = metadata.get("format")
    supported, reason = True, None
    if fmt == "xlsx" and source.get("blob"):
        try:
            _xlsx_patch(doc, source["blob"], metadata)
        except HandoffError as exc:
            supported, reason = False, str(exc)
    return {"format": fmt, "delivery_mode": metadata.get("delivery_mode", "original_patch"),
            "original_patch_supported": supported, "reason": reason,
            "can_export_new_workbook": bool(_tables(doc)), "can_download_source": bool(source.get("blob")),
            "worksheets": [{k: v for k, v in mapping.items() if k != "path"} for mapping in metadata.get("worksheets", [])]}


def _match_nodes(old, new, *, rows=False):
    """Rows only match unique content/keys; position never proves row identity."""
    used, matched = set(), {}
    is_header = lambda n: bool(n.get("content")) and all(c.get("type") == "tableHeader" for c in n["content"])
    for key_fn in ([lambda n: (is_header(n), tuple(node_text(c) for c in n.get("content", []))),
                    lambda n: (is_header(n), node_text((n.get("content") or [{}])[0])) if node_text((n.get("content") or [{}])[0]) else None] if rows else [_identity]):
        old_keys = [key_fn(n) for n in old]
        new_keys = [key_fn(n) for n in new]
        oc, nc = Counter(old_keys), Counter(new_keys)
        for j, key in enumerate(new_keys):
            if j not in matched and key and oc[key] == nc[key] == 1:
                i = old_keys.index(key)
                if i not in used and old[i]["type"] == new[j]["type"]:
                    matched[j], _ = i, used.add(i)
    for j, node in enumerate(new):
        if not rows and node["type"] != "table" and j not in matched and j < len(old) and j not in used and old[j]["type"] == node["type"]:
            matched[j], _ = j, used.add(j)
    return matched


def reconcile_ids(old_doc: dict, new_doc: dict) -> dict:
    doc = copy.deepcopy(new_doc)
    old, new = old_doc.get("content", []), doc.get("content", [])
    for j, i in _match_nodes(old, new).items():
        for key in ("bid", "owner", "created"):
            if key in old[i].get("attrs", {}):
                new[j].setdefault("attrs", {})[key] = old[i]["attrs"][key]
        if new[j]["type"] == "table":
            a, b = old[i].get("content", []), new[j].get("content", [])
            old_ids = old[i].get("attrs", {}).get("column_ids", [])
            old_names = [node_text(c) for c in (a[0].get("content", []) if a else [])]
            new_names = [node_text(c) for c in (b[0].get("content", []) if b else [])]
            # Only exact, unique headers carry identities across a changed schema.
            if old_names == new_names and len(old_ids) == len(new_names):
                new[j]["attrs"]["column_ids"] = old_ids
            for y, x in _match_nodes(a, b, rows=True).items():
                for key in ("bid", "owner", "created"):
                    if key in a[x].get("attrs", {}):
                        b[y].setdefault("attrs", {})[key] = a[x]["attrs"][key]
    return doc


def materialize_fields(doc, values):
    """Comparison/export text must include current field values, not placeholders."""
    result = copy.deepcopy(doc)
    def walk(node):
        kind, attrs = node.get("type"), node.get("attrs", {})
        if kind in ("fill", "status"):
            value = values.get(attrs.get("fid"))
            text = value.get("value", "") if isinstance(value, dict) else value if value is not None else attrs.get("label", "") if kind == "status" else ""
            node.clear()
            node.update({"type": "text", "text": str(text)})
            return
        for child in node.get("content", []):
            walk(child)
    walk(result)
    return result


def replacement_preview(current, incoming, values=None, *, limit=50):
    """Bounded whole-file replacement comparison; deliberately not a merge."""
    before = materialize_fields(current, values or {})
    after = reconcile_ids(before, incoming)
    summary = {"added": 0, "removed": 0, "changed": 0, "unchanged": 0}
    changes, ambiguities = [], []
    total_changes = 0
    def add(kind, rid, label, column, old, new):
        nonlocal total_changes
        total_changes += 1
        if len(changes) < limit:
            changes.append({"kind": kind, "row_id": rid, "label": label, "column": column,
                            "before": old[:2000], "after": new[:2000], "value_truncated": len(old) > 2000 or len(new) > 2000})
    def data_rows(table):
        return [row for row in table.get("content", []) if not (row.get("content") and all(c.get("type") == "tableHeader" for c in row["content"]))]
    old_tables = {t["attrs"]["bid"]: t for t in _tables(before)}
    new_tables = {t["attrs"]["bid"]: t for t in _tables(after)}
    for tid in dict.fromkeys([*old_tables, *new_tables]):
        old, new = old_tables.get(tid, {}), new_tables.get(tid, {})
        old_rows, new_rows = data_rows(old), data_rows(new)
        old_map = {r["attrs"]["bid"]: r for r in old_rows}
        new_map = {r["attrs"]["bid"]: r for r in new_rows}
        name = new.get("attrs", {}).get("worksheet_name") or old.get("attrs", {}).get("worksheet_name") or "表格"
        old_keys = Counter(node_text((row.get("content") or [{}])[0]) for row in old_rows)
        new_keys = Counter(node_text((row.get("content") or [{}])[0]) for row in new_rows)
        for key in old_keys.keys() & new_keys.keys():
            unmatched = any(rid not in new_map and node_text((row.get("content") or [{}])[0]) == key for rid, row in old_map.items())
            if unmatched and (old_keys[key] > 1 or new_keys[key] > 1):
                ambiguities.append({"table": name, "key": key[:160], "old_count": old_keys[key], "new_count": new_keys[key],
                    "message": "重複列名無法唯一對位；不依位置沿用列身分。舊列留言會保留為未附著，新列取得新身分。"})
        old_header = old.get("content", [{}])[0].get("content", []) if old.get("content") else []
        new_header = new.get("content", [{}])[0].get("content", []) if new.get("content") else []
        if old and new and [node_text(c) for c in old_header] != [node_text(c) for c in new_header]:
            add("header", None, name, "表頭／欄位", " | ".join(node_text(c) for c in old_header), " | ".join(node_text(c) for c in new_header))
        for rid in dict.fromkeys([*old_map, *new_map]):
            a, b = old_map.get(rid), new_map.get(rid)
            av = [node_text(c) for c in (a or {}).get("content", [])]
            bv = [node_text(c) for c in (b or {}).get("content", [])]
            label = f"{name} · " + ((av or bv or [rid])[0] or rid)
            if not a or not b:
                kind = "added" if not a else "removed"
                summary[kind] += 1
                add(kind, rid, label, "整列", " | ".join(av), " | ".join(bv))
            elif av == bv:
                summary["unchanged"] += 1
            else:
                summary["changed"] += 1
                for c in range(max(len(av), len(bv))):
                    left, right = av[c] if c < len(av) else "", bv[c] if c < len(bv) else ""
                    if left != right:
                        add("changed", rid, label, node_text(new_header[c]) if c < len(new_header) else f"欄 {c + 1}", left, right)
    # Documents and text outside tables also participate in full replacement.
    old_blocks = [n for n in before.get("content", []) if n.get("type") != "table" and node_text(n)]
    new_blocks = [n for n in after.get("content", []) if n.get("type") != "table" and node_text(n)]
    for i in range(max(len(old_blocks), len(new_blocks))):
        left = node_text(old_blocks[i]) if i < len(old_blocks) else ""
        right = node_text(new_blocks[i]) if i < len(new_blocks) else ""
        kind = "unchanged" if left == right else "added" if not left else "removed" if not right else "changed"
        summary[kind] += 1
        if kind != "unchanged":
            add(kind, None, f"段落 {i + 1}", "文件", left, right)
    return {"mode": "replace_all", "summary": summary, "changes": changes, "truncated": total_changes > limit,
            "change_count": total_changes, "ambiguities": ambiguities[:50],
            "message": "此操作以外部檔案取代目前內容，不會合併本地與外部修改；未出現在新檔的內容會被移除。"}


def _patch_zip(raw, patches):
    output = io.BytesIO()
    with _archive(raw) as source, zipfile.ZipFile(output, "w") as dest:
        dest.comment = source.comment
        for info in source.infolist():
            dest.writestr(info, patches.get(info.filename, source.read(info.filename)))
    return output.getvalue()


def _structure(before, after):
    return ([n.get("attrs", {}).get("bid") for n in before] ==
            [n.get("attrs", {}).get("bid") for n in after] and
            [n["type"] for n in before] == [n["type"] for n in after])


def _xlsx_patch(doc, raw, metadata):
    import openpyxl
    baseline = metadata["baseline"]
    if not _structure(baseline.get("content", []), doc.get("content", [])):
        raise HandoffError("XLSX structural edits cannot be patched safely; export CSV or a new tab instead")
    if any(a != b for a, b in zip(baseline.get("content", []), doc.get("content", []))
           if a["type"] != "table"):
        raise HandoffError("Text edits outside worksheet tables cannot be patched into XLSX")
    before, after = _tables(baseline), _tables(doc)
    patches = {}
    with _archive(raw) as archive:
        for old, new, mapping in zip(before, after, metadata["worksheets"]):
            if column_ids(old) != column_ids(new) or not _structure(old.get("content", []), new.get("content", [])) or any(
                    len(a.get("content", [])) != len(b.get("content", []))
                    for a, b in zip(old.get("content", []), new.get("content", []))):
                raise HandoffError("XLSX structural row/column edits cannot be patched safely")
            a, b = _grid(old), _grid(new)
            if a == b:
                continue
            root = _xml(archive.read(mapping["path"]))
            sheet_data = root.find(f"{{{S}}}sheetData")
            for y, (row_a, row_b) in enumerate(zip(a, b), mapping.get("start_row", 1)):
                for x, (value_a, value_b) in enumerate(zip(row_a, row_b), mapping.get("start_column", 1)):
                    if value_a == value_b:
                        continue
                    coordinate = f"{openpyxl.utils.get_column_letter(x)}{y}"
                    cell = sheet_data.find(f"{{{S}}}row/{{{S}}}c[@r='{coordinate}']")
                    if cell is not None and cell.find(f"{{{S}}}f") is not None:
                        formula = cell.find(f"{{{S}}}f")
                        if formula.get("t") in ("array", "shared", "dataTable"):
                            raise HandoffError("Shared/array formula edits cannot be patched safely")
                    # Merged non-anchor cells are not independently editable.
                    for merged in root.findall(f"{{{S}}}mergeCells/{{{S}}}mergeCell"):
                        bounds = openpyxl.worksheet.cell_range.CellRange(merged.get("ref"))
                        if coordinate in bounds and (x, y) != (bounds.min_col, bounds.min_row):
                            raise HandoffError("Merged-cell edits cannot be patched safely")
                    if cell is None:
                        row_node = sheet_data.find(f"{{{S}}}row[@r='{y}']")
                        if row_node is None:
                            row_node = ET.SubElement(sheet_data, f"{{{S}}}row", r=str(y))
                        cell = ET.Element(f"{{{S}}}c", r=coordinate)
                        index = next((i for i, c in enumerate(row_node)
                                      if openpyxl.utils.cell.coordinate_to_tuple(c.get("r"))[1] > x), len(row_node))
                        row_node.insert(index, cell)
                    for child in list(cell):
                        if child.tag in (f"{{{S}}}f", f"{{{S}}}v", f"{{{S}}}is"):
                            cell.remove(child)
                    cell.attrib.pop("t", None)
                    if value_b.startswith("="):
                        ET.SubElement(cell, f"{{{S}}}f").text = value_b[1:]
                    elif re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?(?:[eE][+-]?\d+)?", value_b):
                        ET.SubElement(cell, f"{{{S}}}v").text = value_b
                    else:
                        cell.set("t", "inlineStr")
                        text = ET.SubElement(ET.SubElement(cell, f"{{{S}}}is"), f"{{{S}}}t")
                        text.set(f"{{{XML}}}space", "preserve")
                        text.text = value_b
            patches[mapping["path"]] = ET.tostring(root, xml_declaration=True, encoding="UTF-8")
        if patches:
            workbook = _xml(archive.read("xl/workbook.xml"))
            calculation = workbook.find(f"{{{S}}}calcPr")
            if calculation is None:
                calculation = ET.SubElement(workbook, f"{{{S}}}calcPr")
            calculation.set("calcMode", "auto")
            calculation.set("fullCalcOnLoad", "1")
            calculation.set("forceFullCalc", "1")
            patches["xl/workbook.xml"] = ET.tostring(workbook, xml_declaration=True, encoding="UTF-8")
    return _patch_zip(raw, patches)


def _docx_patch(doc, raw, metadata):
    old, new = metadata["baseline"].get("content", []), doc.get("content", [])
    if not _structure(old, new):
        raise HandoffError("DOCX structural edits cannot be patched safely; export Markdown or a new tab instead")
    mapped = len(metadata["paragraph_indices"])
    if any(a != b for a, b in zip(old[mapped:], new[mapped:])):
        raise HandoffError("DOCX edits outside imported paragraphs cannot be patched safely; export Markdown instead")
    with _archive(raw) as archive:
        root = _xml(archive.read("word/document.xml"))
    body = root.find(f"{{{W}}}body")
    def patch_text(paragraph, value):
        if any(paragraph.find(f".//{{{W}}}{kind}") is not None
               for kind in ("fldChar", "instrText", "drawing", "tab", "br", "sym")):
            raise HandoffError("DOCX fields, drawings, tabs or breaks in edited paragraphs cannot be patched safely")
        texts = paragraph.findall(f".//{{{W}}}t")
        if not texts:
            container = paragraph if paragraph.tag == f"{{{W}}}p" else paragraph.find(f"{{{W}}}p")
            if container is None:
                container = ET.SubElement(paragraph, f"{{{W}}}p")
            texts = [ET.SubElement(ET.SubElement(container, f"{{{W}}}r"), f"{{{W}}}t")]
        offset = 0
        for n, text in enumerate(texts):
            length = len(text.text or "") if n < len(texts) - 1 else len(value) - offset
            text.text = value[offset:offset + length]
            text.set(f"{{{XML}}}space", "preserve")
            offset += length
    for a, b, index in zip(old, new, metadata["paragraph_indices"]):
        if a.get("type") == "table":
            ga, gb = _grid(a), _grid(b)
            if len(ga) != len(gb) or any(len(x) != len(y) for x, y in zip(ga, gb)):
                raise HandoffError("DOCX table structure cannot be patched safely")
            rows = body[index].findall(f"{{{W}}}tr")
            for y, (ra, rb) in enumerate(zip(ga, gb)):
                cells = rows[y].findall(f"{{{W}}}tc")
                for x, (va, vb) in enumerate(zip(ra, rb)):
                    if va != vb:
                        if cells[x].find(f".//{{{W}}}tbl") is not None:
                            raise HandoffError("Nested DOCX tables cannot be patched safely")
                        patch_text(cells[x], vb)
        elif node_text(a) != node_text(b):
            patch_text(body[index], node_text(b))
    return _patch_zip(raw, {"word/document.xml": ET.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)})


def _markdown(doc):
    lines = []
    for node in doc.get("content", []):
        kind = node["type"]
        if kind == "table":
            for i, row in enumerate(_grid(node)):
                lines.append("| " + " | ".join(v.replace("|", "\\|").replace("\n", " ") for v in row) + " |")
                if i == 0:
                    lines.append("| " + " | ".join("---" for _ in row) + " |")
        elif kind == "heading":
            lines.append("#" * node.get("attrs", {}).get("level", 1) + " " + node_text(node))
        elif kind in ("bulletList", "orderedList", "taskList"):
            lines.extend("- " + node_text(item) for item in node.get("content", []))
        else:
            lines.append(node_text(node))
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def _export_grid(table):
    """Native canvas references count data rows; exported Office includes its header."""
    grid = _grid(table)
    has_header = bool(table.get("content") and all(c.get("type") == "tableHeader" for c in table["content"][0].get("content", [])))
    if table.get("attrs", {}).get("worksheet_name") or not has_header:
        return grid
    from openpyxl.formula.tokenizer import Tokenizer, TokenizerError
    def translate(value):
        if not value.startswith("="):
            return value
        try:
            tokens = Tokenizer(value).items
            for token in tokens:
                if token.type == "OPERAND" and token.subtype == "RANGE" and "!" not in token.value:
                    token.value = re.sub(r"(\$?[A-Za-z]{1,3}\$?)(\d+)", lambda m: m[1] + str(int(m[2]) + 1), token.value)
            return "=" + "".join(t.value for t in tokens)
        except TokenizerError:
            raise HandoffError("Invalid formula cannot be exported safely") from None
    return [grid[0], *[[translate(v) for v in row] for row in grid[1:]]]


def export_file(doc: dict, fmt: str, *, source_blob: bytes | None = None, source_meta: dict | None = None, mode="original_patch") -> bytes:
    fmt = fmt.lower().lstrip(".")
    if fmt not in MIME:
        raise HandoffError("Supported export formats: xlsx, docx, csv, md")
    metadata = source_meta or {}
    if mode not in ("original_patch", "new_workbook", "source"):
        raise HandoffError("Invalid export mode")
    if mode == "source":
        if not source_blob:
            raise HandoffError("No original source file")
        return bytes(source_blob)
    if source_blob and metadata.get("format") == fmt and mode != "new_workbook":
        if doc == metadata.get("baseline"):
            return bytes(source_blob)
        if fmt == "xlsx":
            return _xlsx_patch(doc, bytes(source_blob), metadata)
        if fmt == "docx":
            return _docx_patch(doc, bytes(source_blob), metadata)
    if fmt == "md":
        return _markdown(doc).encode("utf-8")
    if fmt == "csv":
        tables = _tables(doc)
        if len(tables) != 1:
            raise HandoffError("CSV export requires exactly one worksheet/table")
        output = io.StringIO(newline="")
        csv.writer(output).writerows(_export_grid(tables[0]))
        return output.getvalue().encode("utf-8-sig")
    output = io.BytesIO()
    if fmt == "xlsx":
        import openpyxl
        book = openpyxl.Workbook()
        book.remove(book.active)
        for index, table in enumerate(_tables(doc), 1):
            name = table.get("attrs", {}).get("worksheet_name") or f"Sheet{index}"
            sheet = book.create_sheet(re.sub(r"[\\/*?:\[\]]", "_", name)[:31])
            types = table.get("attrs", {}).get("column_types", [])
            # Keep original coordinates when the user selected a later header.
            # Prefix values are retained as plain cells, without claiming format fidelity.
            start_row = table.get("attrs", {}).get("source_header_row", 1)
            if start_row > 1 and source_blob and metadata.get("format") == "xlsx":
                source_book = openpyxl.load_workbook(io.BytesIO(source_blob), read_only=True, data_only=False)
                try:
                    if name in source_book.sheetnames:
                        for prefix in source_book[name].iter_rows(max_row=start_row - 1, values_only=True):
                            sheet.append(list(prefix))
                finally:
                    source_book.close()
            elif start_row > 1:
                for _ in range(start_row - 1):
                    sheet.append([None])
            for y, row in enumerate(_export_grid(table)):
                typed = []
                for x, value in enumerate(row):
                    if y > 0 and x < len(types) and types[x] == "number" and re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?(?:[eE][+-]?\d+)?", value):
                        value = float(value) if any(c in value.lower() for c in ".e") else int(value)
                    typed.append(value)
                sheet.append(typed)
        if not book.worksheets:
            book.create_sheet("Sheet1")
        book.save(output)
    else:
        from docx import Document
        document = Document()
        for node in doc.get("content", []):
            if node["type"] == "table":
                grid = _grid(node)
                table = document.add_table(rows=0, cols=max((len(r) for r in grid), default=1))
                for row in grid:
                    for cell, value in zip(table.add_row().cells, row):
                        cell.text = value
            elif node["type"] == "heading":
                document.add_heading(node_text(node), level=node.get("attrs", {}).get("level", 1))
            else:
                document.add_paragraph(node_text(node))
        document.save(output)
    return output.getvalue()
