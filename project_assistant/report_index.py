"""行銷 Excel 報表 → 只有結構的索引（給專案助理看）。數字在這一步就擋掉，外洩檢查通過才發布。

索引裡有：工作表、表頭、每一欄是手動輸入還是公式、公式寫法（常數照 masking 規則遮掉）、
日期範圍與每一欄最後有填的日期、文字格（照 masking 規則遮掉）、Power Query 程式（長字串省略）。
沒有：任何手動輸入的數值、公式算出來的值。
"""
from __future__ import annotations

import base64
import collections
import datetime as dt
import io
import json
import re
import shutil
import threading
import uuid
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from project_assistant.masking import keep_number, mask_formula, mask_text

LOCK = threading.RLock()  # 同一個行程裡，讀索引檔和整包替換不能同時發生（Windows 改名會撞到開著的檔）

SHORT = [  # 檔名關鍵字 → 給助理用的代號與用途（用途照專案頁「現行報表研究」）
    ("每日廣告花費成效", "daily", "每日廣告花費成效：每天彙整各渠道的成效，裡面有一區是「每日提供給營運的數據」"),
    ("每日Google自操", "google", "每日 Google 自操廣告成效：Google 各版本活動每天的成效"),
    ("自操廣告ROAS", "roas", "自操廣告 ROAS：各渠道按安裝日，追蹤首日到第 90 天的付費"),
    ("Meta廣告成效", "meta", "Meta 廣告成效：Meta 廣告組合和素材的分析，可以篩選、比較不同期間"),
    ("週執行成效", "weekly", "週執行成效：週報和簡報用的數字"),
]


ERRORS = {"#DIV/0!", "#VALUE!", "#REF!", "#N/A", "#NAME?", "#NUM!", "#NULL!"}


HEADER_WORDS = {"日期", "天數", "Date", "date", "DATE"}


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        m = re.match(r"^\s*(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", v)
        if m:
            try:
                return dt.date(int(m[1]), int(m[2]), int(m[3]))
            except ValueError:
                return None
    return None


def ranges(rows):
    """[5,6,7,9] → [[5,7],[9,9]]"""
    out = []
    for r in sorted(set(rows)):
        if out and r == out[-1][1] + 1:
            out[-1][1] = r
        else:
            out.append([r, r])
    return out


def package_info(path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
    return {
        "外部連結": sum(1 for n in names if n.startswith("xl/externalLinks/externalLink")),
        "資料連線": "xl/connections.xml" in names,
        "樞紐分析表": sum(1 for n in names if n.startswith("xl/pivotTables/pivotTable")),
        "圖表": sum(1 for n in names if n.startswith("xl/charts/chart")),
    }


def power_query(path):
    out = []
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not name.startswith("customXml/item") or name.endswith(".rels"):
                continue
            blob = z.read(name)
            text = blob.decode("utf-16" if blob[:2] in (b"\xff\xfe", b"\xfe\xff") else "utf-8", "replace")
            m = re.search(r"<DataMashup[^>]*>([^<]+)</DataMashup>", text)
            if not m:
                continue
            raw = base64.b64decode(m[1])
            size = int.from_bytes(raw[4:8], "little")
            pkg = zipfile.ZipFile(io.BytesIO(raw[8:8 + size]))
            for part in pkg.namelist():
                if part.endswith(".m"):
                    code = pkg.read(part).decode("utf-8-sig", "replace")
                    # 「輸入資料」這類表會把整張表壓成長字串嵌在程式裡，一律省略
                    code = re.sub(r'"[^"\n]{60,}"', '"<長字串省略>"', code)
                    masked = "\n".join(mask_text(line)[0] for line in code.splitlines())
                    out.append({"part": part, "code": masked})
    return out


def header_row(ws):
    for r in range(1, min(ws.max_row, 10) + 1):
        for c in range(1, min(ws.max_column, 14) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip() in HEADER_WORDS:
                return r, c
    for r in range(1, min(ws.max_row, 6) + 1):
        texts = [c for c in range(1, ws.max_column + 1) if isinstance(ws.cell(r, c).value, str)]
        if len(texts) >= 2:
            return r, None
    return None, None


def group_labels(ws, row):
    """表頭上一列的群組名稱（有合併儲存格就照合併範圍，沒有就往右延伸）。"""
    groups = {}
    if not row or row < 1:
        return groups
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row:
            v = ws.cell(rng.min_row, rng.min_col).value
            if isinstance(v, str) and v.strip() and not v.startswith("="):
                for c in range(rng.min_col, rng.max_col + 1):
                    groups[c] = mask_text(re.sub(r"\s+", " ", v.strip()))[0][:40]
    last = ""
    for c in range(1, ws.max_column + 1):
        if c in groups:
            last = groups[c]
            continue
        v = ws.cell(row, c).value
        if isinstance(v, str) and v.strip() and not v.startswith("="):
            last = mask_text(re.sub(r"\s+", " ", v.strip()))[0][:40]
        if last:
            groups[c] = last
    return groups


def index_sheet(ws_f, ws_v):
    hr, dc = header_row(ws_f)
    groups = group_labels(ws_f, hr - 1) if hr and hr > 1 else {}
    labels = {}
    if hr:
        for c in range(1, ws_f.max_column + 1):
            v = ws_f.cell(hr, c).value
            base = mask_text(re.sub(r"\s+", " ", v.strip()))[0][:60] if isinstance(v, str) and not v.startswith("=") else ""
            g = groups.get(c, "")
            labels[c] = (g + "／" + base) if g and base and g != base else (base or g)
    dates = {}
    if dc:
        for r in range((hr or 0) + 1, ws_f.max_row + 1):
            d = as_date(ws_v.cell(r, dc).value) or as_date(ws_f.cell(r, dc).value)
            if d:
                dates[r] = d
    cols = collections.defaultdict(lambda: {
        "formula_rows": collections.defaultdict(list), "manual_rows": [], "text": collections.Counter(),
        "date_rows": [], "errors": 0, "masked_consts": 0,
    })
    texts_all = []
    refs_out = collections.Counter()
    for row in ws_f.iter_rows():
        for cell in row:
            v = cell.value
            if v is None:
                continue
            c = cell.column
            col = cols[c]
            if isinstance(v, ArrayFormula):
                v = v.text or ""
            if isinstance(v, str) and v.startswith("="):
                pattern, masked = mask_formula(v, cell.row)
                col["formula_rows"][pattern].append(cell.row)
                col["masked_consts"] += masked
                for m in re.finditer(r"(?:'([^']+)'|([^\s'!=(),+\-*/&<>:;\"]+))!", v):
                    other = m[1] or m[2]
                    if other != ws_f.title:
                        refs_out[other] += 1
            elif isinstance(v, bool):
                col["text"][str(v)] += 1
            elif isinstance(v, (int, float)):
                col["manual_rows"].append(cell.row)
            elif isinstance(v, (dt.datetime, dt.date)):
                col["date_rows"].append(cell.row)
            elif isinstance(v, str):
                text = mask_text(re.sub(r"\s+", " ", v.strip()))[0][:80]
                if text:
                    col["text"][text] += 1
                    texts_all.append([cell.coordinate, text])
    for row in ws_v.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() in ERRORS:
                cols[cell.column]["errors"] += 1
    columns = {}
    for c in sorted(cols):
        col = cols[c]
        letter = get_column_letter(c)
        entry = {"label": labels.get(c, "")}
        if col["formula_rows"]:
            pats = sorted(col["formula_rows"].items(), key=lambda kv: -len(kv[1]))
            entry["formulas"] = [{"pattern": p, "rows": ranges(rs), "n": len(rs)} for p, rs in pats]
            if col["masked_consts"]:
                entry["formula_hand_typed_numbers"] = col["masked_consts"]
        if col["manual_rows"]:
            filled = [dates[r] for r in col["manual_rows"] if r in dates]
            entry["manual"] = {"n": len(col["manual_rows"]), "rows": ranges(col["manual_rows"]),
                               "last_date": max(filled).isoformat() if filled else None,
                               "first_date": min(filled).isoformat() if filled else None}
        if col["text"]:
            distinct = col["text"]
            entry["texts"] = {"distinct": len(distinct), "top": [[t, n] for t, n in distinct.most_common(12)]}
        if col["date_rows"]:
            entry["date_cells"] = len(col["date_rows"])
        if col["errors"]:
            entry["error_cells"] = col["errors"]
        columns[letter] = entry
    manual_any = set()
    for c in cols:
        if c != dc:
            manual_any.update(cols[c]["manual_rows"])
    filled_rows = [r for r in dates if r in manual_any]
    sheet = {
        "name": ws_f.title,
        "hidden": ws_f.sheet_state != "visible",
        "rows": ws_f.max_row,
        "cols": ws_f.max_column,
        "header_row": hr,
        "date_col": get_column_letter(dc) if dc else None,
        "date_range": [min(dates.values()).isoformat(), max(dates.values()).isoformat()] if dates else None,
        "last_filled_date": max(dates[r] for r in filled_rows).isoformat() if filled_rows else None,
        "titles": [t for ref, t in texts_all if hr and int(re.sub(r"[A-Z]+", "", ref)) < hr][:12],
        "columns": columns,
        "texts_all": texts_all[:4000],
        "refs_out": dict(refs_out.most_common(20)),
        "dates_by_row": {str(r): d.isoformat() for r, d in sorted(dates.items())},
    }
    return sheet


def short_name(path):
    flat = path.name.replace(" ", "")
    for key, code, purpose in SHORT:
        if key in flat:
            return code, purpose
    return path.stem, ""


# ---------- 外洩檢查（測試版 leak_check.py） ----------

TEXT_NUM = re.compile(r"(?<![A-Za-z0-9_\-.])(\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)(%?)(?![A-Za-z0-9_\-]|\.\d)")


DATE_LIKE = re.compile(r"\d{4}[/\-.]\d{1,2}[/\-.]\d{1,2}|\d{1,2}[/\-]\d{1,2}(?:[/\-]\d{2,4})?|\d{1,2}:\d{2}")


def text_violations(text):
    work = DATE_LIKE.sub(" ", str(text))
    bad = 0
    for m in TEXT_NUM.finditer(work):
        number, percent = m.group(1), m.group(2)
        if percent:
            if "." in number or "," in number or float(number.replace(",", "")) > 100:
                bad += 1
        elif not keep_number(number):
            bad += 1
    return bad


def formula_violations(pattern):
    work = str(pattern)
    bad = 0
    for lit in re.findall(r'"([^"]*)"', work):
        bad += text_violations(lit)
    work = re.sub(r'"[^"]*"', " ", work)
    work = re.sub(r"'[^']*'!", " ", work)  # 有引號的工作表名稱
    work = re.sub(r"[^\s'!=(),+\-*/&<>:;\"]+!", " ", work)  # 沒引號的工作表名稱
    work = re.sub(r"\[r(?:[+-]\d+)?\]", " ", work)  # 相對列號
    work = re.sub(r"(?<![A-Za-z0-9_])\$?[A-Z]{1,3}\$?\d+", " ", work)  # 絕對參照
    work = re.sub(r"[A-Z][A-Z0-9.]*\(", "(", work)  # 函數名稱（例如 LOG10）
    for m in re.finditer(r"(?<![A-Za-z0-9_.])\d+(?:\.\d+)?(?![A-Za-z0-9_])", work):
        if not keep_number(m.group(0)):
            bad += 1
    return bad


def walk_book(book, report):
    def check(kind, where, value):
        n = text_violations(value) if kind == "text" else formula_violations(value)
        if n:
            report.append((where, kind, n))

    check("text", f"{book['code']}/purpose", book.get("purpose", ""))
    for i, pq in enumerate(book.get("power_query", [])):
        check("text", f"{book['code']}/power_query[{i}]", pq.get("code", ""))
    for sh in book["sheets"]:
        base = f"{book['code']}/{sh['name']}"
        for t in sh.get("titles", []):
            check("text", base + "/titles", t)
        for ref, t in sh.get("texts_all", []):
            check("text", f"{base}/{ref}", t)
        for other in sh.get("refs_out", {}):
            check("text", base + "/refs_out", other)
        for letter, col in sh.get("columns", {}).items():
            check("text", f"{base}/{letter}/label", col.get("label", ""))
            for t, _ in col.get("texts", {}).get("top", []):
                check("text", f"{base}/{letter}/texts", t)
            for f in col.get("formulas", []):
                check("formula", f"{base}/{letter}/formula", f["pattern"])


# ---------- 建索引、發布 ----------

def build_index(source_dir: Path, out_dir: Path) -> list[dict]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    catalog, used = [], set()
    for path in sorted(Path(source_dir).glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        code, purpose = short_name(path)
        base, n = code, 2
        while code in used:
            code, n = f"{base}-{n}", n + 1
        used.add(code)
        wb_f = openpyxl.load_workbook(path, data_only=False)
        wb_v = openpyxl.load_workbook(path, data_only=True)
        try:
            book = {
                "code": code,
                "file": path.name,
                "purpose": purpose,
                "package": package_info(path),
                "power_query": power_query(path),
                "sheets": [index_sheet(ws, wb_v[ws.title]) for ws in wb_f.worksheets],
            }
        finally:
            wb_f.close()
            wb_v.close()
        (out_dir / f"{code}.json").write_text(json.dumps(book, ensure_ascii=False), encoding="utf-8")
        catalog.append({"code": code, "file": path.name, "purpose": purpose,
                        "sheets": [s["name"] for s in book["sheets"]]})
    (out_dir / "catalog.json").write_text(json.dumps(catalog, ensure_ascii=False, indent=1), encoding="utf-8")
    return catalog


def leak_report(index_dir: Path) -> list[tuple[str, str, int]]:
    report: list[tuple[str, str, int]] = []
    for path in sorted(Path(index_dir).glob("*.json")):
        if path.name in ("catalog.json", "manifest.json"):
            continue
        walk_book(json.loads(path.read_text(encoding="utf-8")), report)
    return report


def source_manifest(source_dir: Path) -> list[dict]:
    out = []
    for path in sorted(Path(source_dir).glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        st = path.stat()
        out.append({"name": path.name, "size": st.st_size, "mtime": int(st.st_mtime)})
    return out


def refresh(source_dir, report_dir, *, ensure_local=None) -> dict:
    """來源資料夾有變才重建；外洩檢查沒過就不發布（舊的索引留著）。"""

    source, target = Path(source_dir), Path(report_dir)
    if not source.is_dir():
        return {"state": "no_source", "source": str(source)}
    manifest = source_manifest(source)
    if not manifest:
        return {"state": "empty", "source": str(source)}
    current = target / "manifest.json"
    if current.exists():
        try:
            if json.loads(current.read_text(encoding="utf-8")).get("files") == manifest:
                return {"state": "unchanged", "files": len(manifest)}
        except ValueError:
            pass
    if ensure_local is not None:
        for item in manifest:
            ensure_local(source / item["name"])
    target.parent.mkdir(parents=True, exist_ok=True)
    staging = target.parent / f".{target.name}-new-{uuid.uuid4().hex[:8]}"
    try:
        catalog = build_index(source, staging)
        problems = leak_report(staging)
        if problems:
            return {"state": "leak", "count": sum(n for *_, n in problems), "where": [w for w, _, _ in problems[:20]]}
        published_at = dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")
        (staging / "manifest.json").write_text(
            json.dumps({"files": manifest, "published_at": published_at, "books": [c["code"] for c in catalog]},
                       ensure_ascii=False, indent=1), encoding="utf-8")
        with LOCK:
            old = target.parent / f".{target.name}-old-{uuid.uuid4().hex[:8]}"
            if target.exists():
                target.rename(old)
            staging.rename(target)
        shutil.rmtree(old, ignore_errors=True)
        return {"state": "published", "files": len(manifest), "books": [c["code"] for c in catalog],
                "published_at": published_at}
    finally:
        if staging.exists():
            shutil.rmtree(staging, ignore_errors=True)
