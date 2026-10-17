"""File handoff contracts: byte identity, safe Office patches and concurrent imports."""
import copy
import io
import unittest
import zipfile
from unittest import mock

import openpyxl
from docx import Document
from project_docs import service
try:
    from project_docs import file_handoff as files
except ImportError:
    files = None


def workbook_bytes():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Shots"
    sheet.append(["Shot", "Cost", "Total"])
    sheet.append(["S01", 1200, "=B2*2"])
    sheet["B2"].number_format = "#,##0.00"
    sheet["B2"].font = openpyxl.styles.Font(bold=True, color="FF0000")
    book.create_sheet("Other")["A1"] = "Keep me"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def set_text(node, value):
    node["content"] = [{"type": "paragraph", "content": [{"type": "text", "text": value}]}]


class FileHandoffTests(unittest.TestCase):
    def parse(self, name, data):
        self.assertIsNotNone(files, "file_handoff module must exist")
        return files.import_file(name, data, slug="demo")

    def test_xlsx_no_edit_is_original_bytes_and_all_sheets_visible(self):
        raw = workbook_bytes()
        parsed = self.parse("shots.xlsx", raw)
        tables = [n for n in parsed.doc["content"] if n["type"] == "table"]
        self.assertEqual([n["attrs"]["worksheet_name"] for n in tables], ["Shots", "Other"])
        self.assertEqual(files.export_file(parsed.doc, "xlsx", source_blob=raw,
                                         source_meta=parsed.metadata), raw)

    def test_native_numeric_sheet_exports_correct_excel_coordinates_and_types(self):
        from project_docs import templates, markdown_import
        spec = next(s for s in templates.tab_specs("marketing", "Test", "") if s[0] == "daily-sheet")
        doc, _ = markdown_import.markdown_to_doc(spec[3], tab_keys={"daily-sheet"}, image_src=lambda _: None)
        templates.configure_doc("daily-sheet", doc)
        book = openpyxl.load_workbook(io.BytesIO(files.export_file(doc, "xlsx")))
        self.assertEqual(book.active["E2"].value, "=D2/C2")
        self.assertEqual(book.active["E3"].value, "=D3/C3")
        self.assertEqual(book.active["C2"].value, 12000)
        self.assertEqual(book.active["C2"].data_type, "n")
        self.assertIn("=D2/C2", files.export_file(doc, "csv").decode("utf-8-sig"))

    def test_xlsx_edit_preserves_other_zip_parts_styles_and_formula(self):
        raw = workbook_bytes()
        parsed = self.parse("shots.xlsx", raw)
        table = next(n for n in parsed.doc["content"] if n["type"] == "table")
        set_text(table["content"][1]["content"][1], "2400")
        output = files.export_file(parsed.doc, "xlsx", source_blob=raw, source_meta=parsed.metadata)
        book = openpyxl.load_workbook(io.BytesIO(output))
        self.assertEqual(book["Shots"]["B2"].value, 2400)
        self.assertTrue(book["Shots"]["B2"].font.bold)
        self.assertEqual(book["Shots"]["B2"].number_format, "#,##0.00")
        self.assertEqual(book["Shots"]["C2"].value, "=B2*2")
        with zipfile.ZipFile(io.BytesIO(raw)) as before, zipfile.ZipFile(io.BytesIO(output)) as after:
            changed = [n for n in before.namelist() if before.read(n) != after.read(n)]
        self.assertEqual(set(changed), {"xl/worksheets/sheet1.xml", "xl/workbook.xml"})
        self.assertEqual(book.calculation.calcMode, "auto")
        self.assertTrue(book.calculation.forceFullCalc)

    def test_xlsx_structural_changes_rejected(self):
        raw = workbook_bytes()
        parsed = self.parse("shots.xlsx", raw)
        table = next(n for n in parsed.doc["content"] if n["type"] == "table")
        table["content"].pop()
        with self.assertRaisesRegex(files.HandoffError, "structural"):
            files.export_file(parsed.doc, "xlsx", source_blob=raw, source_meta=parsed.metadata)

    def test_xlsx_merged_cell_and_unmapped_text_edits_rejected(self):
        book = openpyxl.Workbook()
        book.active.append(["Heading", "Value"])
        book.active.merge_cells("A2:B2")
        book.active["A2"] = "Merged"
        buf = io.BytesIO()
        book.save(buf)
        raw = buf.getvalue()
        parsed = self.parse("merged.xlsx", raw)
        set_text(parsed.doc["content"][0]["content"][1]["content"][1], "Unsafe")
        with self.assertRaisesRegex(files.HandoffError, "Merged"):
            files.export_file(parsed.doc, "xlsx", source_blob=raw, source_meta=parsed.metadata)
        parsed = self.parse("merged.xlsx", raw)
        parsed.doc["content"][-1]["content"] = [{"type": "text", "text": "Important note"}]
        with self.assertRaisesRegex(files.HandoffError, "outside"):
            files.export_file(parsed.doc, "xlsx", source_blob=raw, source_meta=parsed.metadata)

    def test_docx_field_edits_rejected_instead_of_corrupting_fields(self):
        doc = Document()
        paragraph = doc.add_paragraph("Field value")
        from docx.oxml import OxmlElement
        paragraph.runs[0]._r.append(OxmlElement("w:fldChar"))
        buf = io.BytesIO()
        doc.save(buf)
        parsed = self.parse("field.docx", buf.getvalue())
        parsed.doc["content"][0]["content"][0]["text"] = "Changed"
        with self.assertRaisesRegex(files.HandoffError, "fields"):
            files.export_file(parsed.doc, "docx", source_blob=buf.getvalue(), source_meta=parsed.metadata)

    def test_docx_patch_preserves_runs_and_untouched_parts(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Hello ").bold = True
        p.add_run("world").italic = True
        doc.add_paragraph("Untouched")
        buf = io.BytesIO()
        doc.save(buf)
        raw = buf.getvalue()
        parsed = self.parse("script.docx", raw)
        self.assertEqual(files.export_file(parsed.doc, "docx", source_blob=raw,
                                         source_meta=parsed.metadata), raw)
        parsed.doc["content"][0]["content"][0]["text"] = "Goodbye world"
        output = files.export_file(parsed.doc, "docx", source_blob=raw, source_meta=parsed.metadata)
        result = Document(io.BytesIO(output))
        self.assertEqual(result.paragraphs[0].text, "Goodbye world")
        self.assertTrue(result.paragraphs[0].runs[0].bold)
        self.assertTrue(result.paragraphs[0].runs[1].italic)
        with zipfile.ZipFile(io.BytesIO(raw)) as before, zipfile.ZipFile(io.BytesIO(output)) as after:
            self.assertEqual([n for n in before.namelist() if before.read(n) != after.read(n)],
                             ["word/document.xml"])

    def test_table_only_docx_cell_patch_and_unmapped_text_rejection(self):
        doc = Document()
        doc.add_table(rows=2, cols=1).cell(0, 0).text = "Heading"
        doc.tables[0].cell(1, 0).text = "Old"
        buf = io.BytesIO()
        doc.save(buf)
        raw = buf.getvalue()
        parsed = self.parse("table.docx", raw)
        self.assertEqual(files.export_file(parsed.doc, "docx", source_blob=raw, source_meta=parsed.metadata), raw)
        set_text(parsed.doc["content"][0]["content"][1]["content"][0], "New")
        output = files.export_file(parsed.doc, "docx", source_blob=raw, source_meta=parsed.metadata)
        self.assertEqual(Document(io.BytesIO(output)).tables[0].cell(1, 0).text, "New")
        parsed.doc["content"][-1]["content"] = [{"type": "text", "text": "Unmapped note"}]
        with self.assertRaisesRegex(files.HandoffError, "outside"):
            files.export_file(parsed.doc, "docx", source_blob=raw, source_meta=parsed.metadata)

    def test_csv_and_markdown_no_edit_and_edited_export(self):
        for name, raw, fmt in [("data.csv", b'Name,Value\r\n"A,B",1200\r\n', "csv"),
                               ("notes.md", b"# Heading\n\nHello\n", "md")]:
            with self.subTest(name=name):
                parsed = self.parse(name, raw)
                self.assertEqual(files.export_file(parsed.doc, fmt, source_blob=raw,
                                                 source_meta=parsed.metadata), raw)
                output = files.export_file(parsed.doc, "docx")
                self.assertGreater(len(Document(io.BytesIO(output)).paragraphs) + len(Document(io.BytesIO(output)).tables), 0)

    def test_reimport_keeps_row_id_after_insertion_and_edit(self):
        old = self.parse("rows.csv", b"ID,Text\nA,first\nB,second\n")
        new = self.parse("rows.csv", b"ID,Text\nX,new\nA,changed\nB,second\n")
        aligned = files.reconcile_ids(old.doc, new.doc)
        a = old.doc["content"][0]["content"][1]["attrs"]["bid"]
        b = old.doc["content"][0]["content"][2]["attrs"]["bid"]
        self.assertEqual(aligned["content"][0]["content"][2]["attrs"]["bid"], a)
        self.assertEqual(aligned["content"][0]["content"][3]["attrs"]["bid"], b)

    def test_reimport_duplicate_keys_only_preserve_uniquely_identified_rows(self):
        old = self.parse("rows.csv", b"ID,Text\nA,one\nA,two\nB,unique\n")
        reordered = self.parse("rows.csv", b"ID,Text\nA,two\nB,unique\nA,one\n")
        aligned = files.reconcile_ids(old.doc, reordered.doc)
        original_rows = old.doc["content"][0]["content"]
        self.assertEqual([r["attrs"]["bid"] for r in aligned["content"][0]["content"][1:]],
                         [original_rows[i]["attrs"]["bid"] for i in (2, 3, 1)])
        ambiguous = self.parse("rows.csv", b"ID,Text\nA,changed-two\nB,unique\nA,changed-one\n")
        aligned = files.reconcile_ids(old.doc, ambiguous.doc)
        old_duplicate_ids = {original_rows[i]["attrs"]["bid"] for i in (1, 2)}
        self.assertTrue(old_duplicate_ids.isdisjoint({r["attrs"]["bid"] for r in aligned["content"][0]["content"]}))
        self.assertTrue(files.replacement_preview(old.doc, ambiguous.doc)["ambiguities"])

    def test_replacement_preview_uses_current_field_values_and_limits_diff(self):
        old = self.parse("rows.csv", b"ID,Text\nA,old\n")
        old.doc["content"][0]["content"][1]["content"][1]["content"] = [{"type": "fill", "attrs": {"fid": "fill001", "label": "placeholder"}}]
        new = self.parse("rows.csv", b"ID,Text\nA,incoming\nB,added\n")
        preview = files.replacement_preview(old.doc, new.doc, {"fill001": {"value": "local changes"}}, limit=1)
        self.assertEqual(preview["summary"], {"added": 1, "removed": 0, "changed": 1, "unchanged": 0})
        self.assertEqual(preview["changes"][0]["before"], "local changes")
        self.assertEqual(preview["changes"][0]["after"], "incoming")
        self.assertTrue(preview["truncated"])

    def test_malformed_and_expanding_archives_are_rejected(self):
        self.assertIsNotNone(files)
        for name, raw in [("bad.xlsx", b"not zip"), ("bad.md", b"\xff"),
                          ("bad.exe", b"anything")]:
            with self.subTest(name=name), self.assertRaises(files.HandoffError):
                files.import_file(name, raw, slug="demo")
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/workbook.xml", b"a" * 2_000_000)
        with self.assertRaises(files.HandoffError):
            files.import_file("bomb.xlsx", buf.getvalue(), slug="demo")

    def test_import_requires_current_revision_and_checks_locks(self):
        me = service.Identity("11111111-1111-4111-8111-111111111111", "Admin", "admin")
        parsed = self.parse("notes.md", b"Old\n")
        tab = {"doc": parsed.doc, "revision": 4, "title": "Notes", "surface_type": "doc"}
        with mock.patch.object(service.repo, "get_project", return_value={"slug": "demo"}), \
             mock.patch.object(service.repo, "get_tab_for_update", return_value=tab), \
             mock.patch.object(service.repo, "tab_locks", return_value={}):
            status, _ = service.import_tab_file(None, me, "demo", "notes.md", b"New\n", tab_key="notes", base_revision=3)
            self.assertEqual(status, 409)
            status, _ = service.import_tab_file(None, me, "demo", "notes.md", b"New\n", tab_key="notes")
            self.assertEqual(status, 400)
        bid = parsed.doc["content"][0]["attrs"]["bid"]
        with mock.patch.object(service.repo, "get_project", return_value={"slug": "demo"}), \
             mock.patch.object(service.repo, "get_tab_for_update", return_value=tab), \
             mock.patch.object(service.repo, "tab_locks", return_value={bid: "other"}):
            status, _ = service.import_tab_file(None, me, "demo", "notes.md", b"New\n", tab_key="notes", base_revision=4)
            self.assertEqual(status, 409)

    def test_only_admin_can_change_privacy(self):
        member = service.Identity("11111111-1111-4111-8111-111111111111", "User", "marketer")
        with mock.patch.object(service.repo, "get_project", return_value={"slug": "demo"}):
            status, _ = service.update_project(None, member, "demo", {"privacy_mode": "mask_numbers"})
        self.assertEqual(status, 403)

    def test_privacy_rejects_invalid_values_before_writing(self):
        admin = service.Identity("11111111-1111-4111-8111-111111111111", "Admin", "admin")
        with mock.patch.object(service.repo, "get_project", return_value={"slug": "demo"}):
            status, _ = service.update_project(None, admin, "demo", {"privacy_mode": "other"})
        self.assertEqual(status, 400)
        with mock.patch.object(service.repo, "get_project", return_value=None):
            status, _ = service.create_project(None, admin, {"slug": "new", "title": "New", "privacy_mode": "other"})
        self.assertEqual(status, 400)

    def test_import_updates_source_and_reanchors_in_same_transaction(self):
        admin = service.Identity("11111111-1111-4111-8111-111111111111", "Admin", "admin")
        old = self.parse("notes.md", b"Old\n")
        tab = {"doc": old.doc, "revision": 1, "title": "Notes", "surface_type": "doc"}
        state = {"tab": tab}
        def save(conn, slug, key, **kwargs):
            state["tab"] = {**tab, "doc": kwargs["doc"], "revision": 2}
            return {"revision": 2}
        def source(conn, slug, key, **kwargs):
            state["source"] = kwargs
        def reanchor(conn, slug, key, doc, **kwargs):
            state["anchored"] = doc
        with mock.patch.object(service.repo, "get_project", return_value={"slug": "demo"}), \
             mock.patch.object(service.repo, "get_tab_for_update", return_value=tab), \
             mock.patch.object(service.repo, "tab_locks", return_value={}), \
             mock.patch.object(service.repo, "save_tab_doc", side_effect=save), \
             mock.patch.object(service.repo, "set_tab_source", side_effect=source), \
             mock.patch.object(service.repo, "get_tab", side_effect=lambda *a: state["tab"]), \
             mock.patch.object(service.repo, "list_tabs", side_effect=lambda *a: [state["tab"]]), \
             mock.patch.object(service.repo, "field_values", return_value={}), \
             mock.patch.object(service.repo, "add_audit"), \
             mock.patch.object(service, "reanchor_tab", side_effect=reanchor):
            status, payload = service.import_tab_file(None, admin, "demo", "notes.md", b"New\n", tab_key="notes", base_revision=1, confirm_replace=True)
        self.assertEqual(status, 200)
        self.assertEqual(payload["tab"]["revision"], 2)
        self.assertEqual(state["source"]["blob"], b"New\n")
        self.assertEqual(state["source"]["metadata"]["baseline"], state["anchored"])
        self.assertEqual(payload["changes"][0]["change"], "changed")


if __name__ == "__main__":
    unittest.main()
