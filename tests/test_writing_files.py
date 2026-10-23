import base64
import copy
from datetime import datetime, timedelta, timezone
import io
import sys
from types import SimpleNamespace
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from project_docs import writing_files as files, writing_tools as tools
from tests.test_writing_tools_fixtures import InterfaceFixtures, OTHER, doc


class BinaryFixture:
    def __init__(self, body, content_type, *, status=200, headers=None):
        self.body, self.content_type, self.status, self.headers = body, content_type, status, headers or {}


def archive(parts):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as z:
        for key, value in parts.items():
            z.writestr(key, value)
    return output.getvalue()


def docx_xml(content):
    return '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' + content + '</w:body></w:document>'


class ParserTests(unittest.TestCase):
    def test_literal_utf8_bom_whitespace_dialogue_and_no_heading_source_loss(self):
        value = "  我心想自己要離開。 😀\r\nKEEPER:\r\n  Stay.  \r\n\r\n"
        parsed = files.parse_upload("novel.txt", b"\xef\xbb\xbf" + value.encode(), "novel")
        self.assertEqual(parsed["text"], value)
        self.assertEqual(parsed["sections"], [{"sectionId": "sec_1", "title": "Imported manuscript", "start": 0, "end": len(value), "included": True}])
        converted = files.literal_doc(value, "demo")
        self.assertEqual("\n".join(tools.block_text(n) for n in converted["content"]), value)
        self.assertIn("心想自己", parsed["text"])

    def test_purpose_specific_grammars_keep_preamble_and_all_source(self):
        value = "Preface\nChapter 1: Arrival\nA keeper.\nChapter 2: Signal\nA lamp.\n"
        parsed = files.parse_upload("novel.md", value.encode(), "novel")
        self.assertEqual(len(parsed["sections"]), 3)
        self.assertEqual("".join(value[s["start"]:s["end"]] for s in parsed["sections"]), value)
        self.assertEqual(len(files.parse_upload("novel.md", value.encode(), "script")["sections"]), 1)
        self.assertTrue(parsed["warnings"])

    def test_real_docx_heading_styles_linebreaks_tabs_tables_and_loss_warnings(self):
        document = Document()
        document.add_heading("Chapter 1: Arrival", 1)
        p = document.add_paragraph()
        r = p.add_run("繁中 😀")
        r.bold = True
        r.add_break()
        r.add_tab()
        r.add_text("  dialogue  ")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Table text"
        raw = io.BytesIO()
        document.save(raw)
        parsed = files.parse_upload("chapter.docx", raw.getvalue(), "novel")
        self.assertIn("繁中 😀\n\t  dialogue  ", parsed["text"])
        self.assertIn("Table text", parsed["text"])
        self.assertEqual(parsed["sections"][0]["title"], "Chapter 1: Arrival")
        self.assertTrue(any("Tables" in w for w in parsed["warnings"]))
        self.assertTrue(any("Rich text" in w for w in parsed["warnings"]))

    def test_docx_tracked_text_preserved_and_warned(self):
        xml = docx_xml('<w:p><w:r><w:t>Old </w:t></w:r><w:del><w:r><w:delText>deleted</w:delText></w:r></w:del><w:ins><w:r><w:t>new</w:t></w:r></w:ins></w:p>')
        parsed = files.parse_upload("story.docx", archive({"word/document.xml": xml}), "script")
        self.assertEqual(parsed["text"], "Old deletednew")
        self.assertTrue(any("Tracked revisions" in w for w in parsed["warnings"]))

    def test_docx_rejects_traversal_external_entities_macros_and_zip_bombs(self):
        xml = docx_xml('<w:p><w:r><w:t>hello</w:t></w:r></w:p>')
        bad = [
            {"word/document.xml": xml, "../outside": "bad"},
            {"word/document.xml": xml, "word\\..\\outside": "bad"},
            {"word/document.xml": xml, "word/_rels/document.xml.rels": '<Relationships><Relationship TargetMode="External" Target="https://example.invalid/remote"/></Relationships>'},
            {"word/document.xml": '<!DOCTYPE x [<!ENTITY xxe SYSTEM "file:///secret">]>' + xml},
            {"word/document.xml": ('<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE x [<!ENTITY xxe "bad">]>' + xml).encode("utf-16")},
            {"word/document.xml": xml, "word/vbaProject.bin": b"macro"},
            {"word/document.xml": xml, "[Content_Types].xml": '<Types><Override ContentType="application/vnd.ms-word.document.macroEnabled.main+xml"/></Types>'},
            {"word/document.xml": xml, "padding": "0" * 1_100_000},
            {"word/document.xml": xml, **{f"parts/{i}.txt": "x" for i in range(2001)}}]
        with patch("urllib.request.urlopen", side_effect=AssertionError("No fetch permitted")) as outbound:
            for parts in bad:
                with self.subTest(parts=list(parts)[:3]):
                    with self.assertRaises(tools.story.StoryError):
                        files.parse_upload("bad.docx", archive(parts), "novel")
            outbound.assert_not_called()

    def test_hardbreaks_survive_all_manuscript_exports(self):
        manuscript = doc("Keeper")
        manuscript["content"][0]["content"] += [{"type": "hardBreak"}, {"type": "text", "text": "  Stay.  "}]
        self.assertIn("Keeper\n  Stay.  ", files.render_export(manuscript, "txt", title="One").decode())
        self.assertIn("Keeper\n  Stay.  ", files.render_export(manuscript, "md", title="One").decode())
        reopened = Document(io.BytesIO(files.render_export(manuscript, "docx", title="One")))
        self.assertIn("Keeper\n  Stay.  ", "\n".join(p.text for p in reopened.paragraphs))
        self.assertIn("Keeper\n  Stay.  ", files.render_export(manuscript, "print", title="One").decode())

    def test_file_limits_and_unknown_encoding_formats_paths(self):
        for name, raw in [("a.epub", b"zip"), ("a.docm", b"zip"), ("../a.txt", b"hello"), ("C:\\a.txt", b"hello"),
                          ("a\r\nHeader.txt", b"hello"), ("a\u202e.txt", b"hello"), ("a.txt", b"\xff"), ("a.txt", b"a\x00"),
                          ("a.txt", b"a" * 2_000_001), ("a.txt", b"a" * 500_001), ("a.docx", b"x" * (8 * 1024 * 1024 + 1))]:
            with self.subTest(name=name, size=len(raw)):
                with self.assertRaises(tools.story.StoryError):
                    files.parse_upload(name, raw, "novel")

    def test_boundary_corrections_complete_partition_and_explicit_exclusion(self):
        value = "😀第一章\nKeep\nExclude\n"
        boundary = value.index("Exclude")
        corrected = files._correct_sections(value, [{"sectionId": "a", "title": "Keep", "start": 0, "end": boundary},
                                                    {"sectionId": "b", "title": "Remove", "start": boundary, "end": len(value), "included": False}])
        self.assertEqual(corrected[1]["start"], boundary)
        for sections in ([{"sectionId": "a", "title": "Lost", "start": 1, "end": len(value)}],
                         [{"sectionId": "a", "title": "Truncated", "start": 0, "end": 2}],
                         [{"sectionId": "a", "title": "All removed", "start": 0, "end": len(value), "included": False}]):
            with self.assertRaises(tools.story.StoryError):
                files._correct_sections(value, sections)

    def test_real_outputs_reopen_utf8_and_safe_print(self):
        value = "繁中 😀 <script>alert(1)</script>\nKEEPER: Stay.  "
        manuscript = files.literal_doc(value, "demo")
        txt = files.render_export(manuscript, "txt", title="Keeper version")
        self.assertIn(value, txt.decode())
        md = files.render_export(manuscript, "md", title="Keeper version").decode()
        self.assertIn("繁中 😀", md)
        raw = files.render_export(manuscript, "docx", title="Keeper version")
        reopened = Document(io.BytesIO(raw))
        joined = "\n".join(p.text for p in reopened.paragraphs)
        self.assertIn("繁中 😀", joined)
        self.assertIn("KEEPER: Stay.  ", joined)
        printed = files.render_export(manuscript, "print", title="<title>").decode()
        self.assertIn("&lt;script&gt;", printed)
        self.assertNotIn("<script>", printed)
        self.assertNotIn("http://", printed)
        self.assertIn("break-before:page", printed)
        with self.assertRaises(tools.story.StoryError):
            files.render_export(manuscript, "pdf", title="No fake PDF")


class FilesInterfaceTests(InterfaceFixtures, unittest.TestCase):
    def setUp(self):
        super().setUp()
        self.binary_patch = patch.dict(sys.modules, {"project_docs.writing_http": SimpleNamespace(BinaryResponse=BinaryFixture)})
        self.binary_patch.start()
        self.addCleanup(self.binary_patch.stop)

    def test_raw_upload_readonly_preview_and_revised_boundaries(self):
        before = copy.deepcopy(self.store.docs)
        raw = "第一章\nKeep\nDiscard\n".encode()
        preview = files.dispatch(None, self.scope, "POST", "imports", {}, {}, raw=raw,
                                 headers={"X-Filename": "novel.txt", "X-Writing-Purpose": "novel"})
        self.assertEqual(self.store.docs, before)
        value = preview["proposedChanges"]["text"]
        boundary = value.index("Discard")
        revised = files.build_preview(None, self.scope, {"kind": "import", "uploadPreviewId": preview["previewId"], "uploadHash": preview["uploadHash"],
                        "sections": [{"sectionId": "keep", "title": "Chapter", "start": 0, "end": boundary},
                                     {"sectionId": "remove", "title": "Excluded", "start": boundary, "end": len(value), "included": False}]})
        self.assertEqual(revised["proposedChanges"]["text"], value)
        self.assertEqual(revised["proposedChanges"]["excludedRanges"][0]["start"], boundary)
        self.assertEqual(self.store.docs, before)
        other = self.store.scope_for(None, SimpleNamespace(actor_id=OTHER, can_write=True, can_edit=True), self.scope.episode_id)
        with self.assertRaises(tools.story.StoryError):
            files.build_preview(None, other, {"kind": "import", "uploadPreviewId": preview["previewId"], "uploadHash": preview["uploadHash"], "sections": preview["sections"]})

    def test_raw_upload_aliases_and_conflicts_are_not_silently_defaulted(self):
        preview = files.upload(None, self.scope, b"Scene 1: Arrival\nKeeper waits.\n",
                               {"X-File-Name": "script.txt", "X-File-Purpose": "script"})
        self.assertEqual(preview["purpose"], "script")
        multilingual = files.upload(None, self.scope, "繁中原稿".encode(), {"X-Filename": "%E8%8D%89%E7%A8%BF.txt"})
        self.assertEqual(multilingual["filename"], "草稿.txt")
        literal_percent = files.upload(None, self.scope, b"Literal percent filename",
            {"X-Filename": "%E8%8D%89%E7%A8%BF%252F100%25.txt"})
        self.assertEqual(literal_percent["filename"], "草稿%2F100%.txt")
        for filename in ("..%2Fa.txt", "a%0D%0AInjected.txt", "%FF.txt"):
            with self.assertRaises(tools.story.StoryError):
                files.upload(None, self.scope, b"A source", {"X-Filename": filename})
        for headers in ({"X-Filename": "a.txt", "X-File-Name": "b.txt"},
                        {"X-Filename": "a.txt", "X-Writing-Purpose": "novel", "X-File-Purpose": "script"}):
            with self.assertRaises(tools.story.StoryError):
                files.upload(None, self.scope, b"A source", headers)

    def test_import_script_creates_owner_branch_never_mainline(self):
        preview = files.upload(None, self.scope, b"Scene 1: Arrival\nA keeper waits.\n", {"X-Filename": "script.txt", "X-Writing-Purpose": "script"})
        before = copy.deepcopy(self.store.docs)
        result = files.apply_preview(None, self.scope, preview, {"actions": [{"actionId": "import_document"}]})
        self.assertEqual(result["sourceRef"]["kind"], "branch")
        branch = self.store.get_record(None, self.scope, "writing_branch", result["resourceId"])
        self.assertIn("keeper waits", "\n".join(tools.block_text(n) for n in branch["doc"]["content"]))
        self.assertEqual(branch["sourceFile"]["uploadHash"], preview["uploadHash"])
        self.assertEqual(before, self.store.docs)

    def test_script_import_projects_lens_without_inventing_preamble_scene(self):
        preview = files.upload(None, self.scope, "Lens: cinema × suspense\nScene 01: Arrival\n@CAST-01 waits.\n".encode(),
                               {"X-Filename": "script.txt", "X-Writing-Purpose": "script"})
        self.assertTrue(any("preamble" in warning for warning in preview["warnings"]))
        result = files.apply_preview(None, self.scope, preview, {"actions": [{"actionId": "import_document"}]})
        branch = self.store.get_record(None, self.scope, "writing_branch", result["resourceId"])
        self.assertEqual(branch["composition"], {"mediumLens": "cinema", "genreLens": "suspense"})
        self.assertEqual(len(branch["structure"]["scenes"]), 1)
        self.assertIn("Lens: cinema", tools.block_text(branch["doc"]["content"][0]))

    def test_source_import_calls_core_create_only_save_contract(self):
        preview = files.upload(None, self.scope, b"I secretly think of home.\n", {"X-Filename": "novel.txt"})
        with patch("project_docs.writing_service.save_source", side_effect=lambda conn, scope, mode, body: self._source_save(mode, body)):
            result = files.apply_preview(None, self.scope, preview, {"actions": [{"actionId": "import_document"}]})
        self.assertEqual(result["sourceRef"], {"kind": "source", "mode": "novel"})
        self.assertEqual(result["version"], 2)

    def _source_save(self, mode, body):
        self.assertEqual(body["baseVersion"], 0)
        self.assertTrue(body["clientMutationId"].startswith("import-"))
        self.assertIn("secretly think", tools.block_text(body["doc"]["content"][0]))
        return self.store.insert_record(None, self.scope, "writing_source", "wsrc_novel",
                                        {**body, "sourceId": "wsrc_novel", "mode": mode, "docHash": self.store.doc_hash(body["doc"])})

    def test_immutable_export_selected_scope_and_current_auth_before_bytes(self):
        current = self.store.resolve_ref(None, self.scope, {"kind": "working"})
        result = files.create_export(None, self.scope, {"sourceRef": {"kind": "working"}, "expectedHash": current["docHash"], "format": "txt"})
        first = files.export_content(None, self.scope, result["exportId"])
        self.store.docs[self.scope.episode_id]["doc"] = doc("New draft")
        second = files.export_content(None, self.scope, result["exportId"])
        self.assertEqual(first.body, second.body)
        self.assertIn(b"keeper waits", second.body)
        self.assertEqual(second.headers["Cache-Control"], "private, no-store")
        self.assertIn("attachment;", second.headers["Content-Disposition"])
        self.store.active = False
        with self.assertRaises(tools.story.StoryError) as caught:
            files.export_content(None, self.scope, result["exportId"])
        self.assertEqual(caught.exception.status, 403)

    def test_export_rechecks_hash_after_fixed_input_render_before_persist(self):
        current = self.store.resolve_ref(None, self.scope, {"kind": "working"})
        render = files.render_export
        def changed_during_render(*args, **kwargs):
            self.assertFalse(getattr(self.store, "project_locked", False))
            raw = render(*args, **kwargs)
            self.store.docs[self.scope.episode_id]["doc"] = doc("Someone changed this source")
            return raw
        with patch.object(files, "render_export", side_effect=changed_during_render):
            with self.assertRaises(tools.story.StoryError) as caught:
                files.create_export(None, self.scope, {"sourceRef": {"kind": "working"}, "expectedHash": current["docHash"], "format": "txt"})
        self.assertEqual(caught.exception.code, "DRAMA_WRITING_PREVIEW_STALE")
        self.assertFalse(any(kind == "writing_export" for kind, _ in self.store.records))
        self.assertTrue(self.store.project_locked)

    def test_export_expiry_foreign_owner_and_print_csp(self):
        current = self.store.resolve_ref(None, self.scope, {"kind": "working"})
        result = files.create_export(None, self.scope, {"sourceRef": {"kind": "working"}, "expectedHash": current["docHash"], "format": "print"})
        response = files.export_content(None, self.scope, result["exportId"])
        self.assertTrue(response.body.startswith(b"<!doctype html>"))
        self.assertIn("default-src 'none'", response.headers["Content-Security-Policy"])
        record = self.store.records["writing_export", result["exportId"]]
        record["ownerId"] = OTHER
        with self.assertRaises(tools.story.StoryError):
            files.export_content(None, self.scope, result["exportId"])
        record["ownerId"] = self.scope.actor_id
        record["expiresAt"] = (datetime.now(timezone.utc) - timedelta(minutes=1)).isoformat()
        with self.assertRaises(tools.story.StoryError):
            files.export_content(None, self.scope, result["exportId"])

    def test_export_scene_selection_bible_optin_and_bad_hash(self):
        resource = self.store.docs[self.scope.episode_id]
        resource["doc"]["content"].extend(doc("UNSELECTED", bid="b_other")["content"])
        resource["structure"]["scenes"].append({"sceneKey": "wsc_other", "headingBid": "b_other"})
        current = self.store.resolve_ref(None, self.scope, {"kind": "working"})
        body = {"sourceRef": {"kind": "working"}, "expectedHash": current["docHash"], "format": "txt", "sceneKeys": ["wsc_arrival"]}
        note = tools.create_bible(None, self.scope, {"entryKind": "note", "title": "Research", "notes": "EXPLICIT_BIBLE"})
        result = files.create_export(None, self.scope, body)
        plain = files.export_content(None, self.scope, result["exportId"]).body
        self.assertNotIn(b"UNSELECTED", plain)
        self.assertNotIn(b"EXPLICIT_BIBLE", plain)
        result = files.create_export(None, self.scope, {**body, "include": {"bibleIds": [note["bibleId"]]}})
        self.assertIn(b"EXPLICIT_BIBLE", files.export_content(None, self.scope, result["exportId"]).body)
        with self.assertRaises(tools.story.StoryError):
            files.create_export(None, self.scope, {**body, "expectedHash": "old"})
        with self.assertRaises(tools.story.StoryError):
            files.create_export(None, self.scope, {**body, "sceneKeys": ["missing"]})


if __name__ == "__main__":
    unittest.main()
