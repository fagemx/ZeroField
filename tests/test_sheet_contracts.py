"""User operations stay atomic across fields, concurrent edits and file delivery."""
import copy
import io
import os
import unittest

import openpyxl
import test_product_http as http_fixture
from project_docs import file_handoff, repository as repo, table_schema
from project_docs.doc_schema import column_ids, node_text, validate_doc
from project_assistant import table_edits
from test_table_proposals import sample


class SheetContractTests(unittest.TestCase):
    def test_legacy_column_identity_is_repeatable_and_rename_preserves_it(self):
        doc = sample()
        first, _ = validate_doc(doc, slug="demo")
        second, _ = validate_doc(doc, slug="demo")
        self.assertEqual(column_ids(first["content"][0]), column_ids(second["content"][0]))
        table = first["content"][0]
        changed = table_schema.edit_table(table, {"operation": "configure", "column_id": column_ids(table)[1], "name": "Review", "column_type": "status", "options": ["old", "done"]})
        self.assertEqual(column_ids(changed), column_ids(table))
        self.assertEqual(node_text(changed["content"][0]["content"][1]), "Review")
        self.assertEqual(changed["attrs"]["column_options"][column_ids(table)[1]]["options"], ["old", "done"])

    def test_cas_matches_field_and_multiline_visible_values(self):
        doc = sample()
        row = doc["content"][0]["content"][1]
        row["content"][0]["content"] = [{"type": "fill", "attrs": {"fid": "fill001", "label": "Placeholder"}}]
        row["content"][1]["content"] = [{"type": "paragraph", "content": [{"type": "text", "text": "one"}, {"type": "hardBreak"}, {"type": "text", "text": "two"}]}]
        edits = [{"row_id": "row00001", "cells": {"0": "new", "1": "next"}, "expected_cells": {"0": "", "1": "one\ntwo"}}]
        table_edits.check_expected(doc, edits, {})
        with self.assertRaises(table_edits.CellConflict) as caught:
            table_edits.check_expected(doc, edits, {"fill001": {"value": "other"}})
        self.assertEqual(caught.exception.conflicts[0]["current"], "other")

    def test_header_mapping_and_new_workbook_keep_original_coordinates(self):
        book = openpyxl.Workbook()
        book.active.append(["Report"])
        book.active.append(["Generated locally"])
        book.active.append(["Name", "Value"])
        book.active.append(["A", 10])
        stream = io.BytesIO()
        book.save(stream)
        raw = stream.getvalue()
        preview = file_handoff.preview_file("test.xlsx", raw, slug="demo")
        self.assertEqual(preview["worksheets"][0]["rows"][2], ["Name", "Value"])
        imported = file_handoff.import_file("test.xlsx", raw, slug="demo", header_row=3)
        self.assertEqual(imported.metadata["worksheets"][0]["data_start_row"], 4)
        imported.doc["content"][0]["content"][1]["content"][1] = table_edits._cell("20")
        patched = file_handoff.export_file(imported.doc, "xlsx", source_blob=raw, source_meta=imported.metadata)
        result = openpyxl.load_workbook(io.BytesIO(patched))
        self.assertEqual(result.active["B4"].value, 20)
        self.assertEqual(result.active["A1"].value, "Report")
        fresh = file_handoff.export_file(imported.doc, "xlsx", source_blob=raw, source_meta=imported.metadata, mode="new_workbook")
        result = openpyxl.load_workbook(io.BytesIO(fresh))
        self.assertEqual(result.active["B4"].value, "20")
        self.assertEqual(result.active["A1"].value, "Report")
        self.assertEqual(file_handoff.export_file(imported.doc, "xlsx", source_blob=raw, mode="source"), raw)


@unittest.skipUnless(os.environ.get("OPENBENCH_TEST_DB_URI"), "set OPENBENCH_TEST_DB_URI")
class SheetHTTPContracts(unittest.TestCase):
    setUp = http_fixture.ProductHTTPTests.setUp
    tearDown = http_fixture.ProductHTTPTests.tearDown
    request = http_fixture.ProductHTTPTests.request

    def tab(self, key="shots"):
        code, payload, _ = self.request(f"/projects/api/{self.slug}/tabs/{key}")
        self.assertEqual(code, 200, payload)
        return payload["tab"]

    def edit(self, tab, edits, **extra):
        payload = {"base_revision": tab["revision"], "row_edits": [], **extra}
        table = tab["doc"]["content"][0]
        for rid, cells in edits:
            row = next(row for row in table["content"] if row["attrs"]["bid"] == rid)
            payload["row_edits"].append({"row_id": rid, "cells": cells,
                "expected_cells": {col: table_edits.cell_value(row["content"][int(col)], {}) for col in cells},
                "column_ids": {col: column_ids(table)[int(col)] for col in cells}})
        return self.request(f"/projects/api/{self.slug}/tabs/{tab['key']}/table-edit", method="POST", data=payload)

    def test_rebase_unrelated_cell_and_inverse_conflict_is_atomic(self):
        base = self.tab()
        code, first, _ = self.edit(base, [("b_shot1", {"0": "renamed"})])
        self.assertEqual(code, 200, first)
        code, second, _ = self.edit(base, [("b_shot1", {"1": "changed"})])
        self.assertEqual(code, 200, second)
        self.assertEqual(second["cell_patches"][0]["before"], "original")
        current = self.tab()
        self.assertEqual(node_text(current["doc"]["content"][0]["content"][1]["content"][0]), "renamed")
        # A later writer changes the value that undo would restore. No other
        # cell in the same attempted inverse is allowed to sneak through.
        self.assertEqual(self.edit(current, [("b_shot1", {"1": "later"})])[0], 200)
        code, conflict, _ = self.edit(current, [("b_shot1", {"1": "original"}), ("b_shot2", {"1": "should not write"})])
        self.assertEqual(code, 409, conflict)
        self.assertEqual(conflict["conflicts"][0]["current"], "later")
        self.assertEqual(node_text(self.tab()["doc"]["content"][0]["content"][2]["content"][1]), "original")

    def test_explicit_rows_never_include_hidden_rows(self):
        doc = copy.deepcopy(self.doc)
        third = copy.deepcopy(doc["content"][0]["content"][2])
        third["attrs"]["bid"] = "b_shot3"
        doc["content"][0]["content"].append(third)
        repo.save_tab_doc(self.conn, self.slug, "shots", doc=doc, title="Shots", actor_id=self.actor, session_id="fixture", base_revision=None)
        self.conn.commit()
        base = self.tab()
        code, result, _ = self.edit(base, [("b_shot1", {"1": "first"}), ("b_shot3", {"1": "third"})])
        self.assertEqual(code, 200, result)
        self.assertEqual([node_text(r["content"][1]) for r in result["doc"]["content"][0]["content"][1:]], ["first", "original", "third"])

    def test_mixed_field_cas_rejects_every_cell_on_one_conflict(self):
        doc = copy.deepcopy(self.doc)
        doc["content"][0]["content"][1]["content"][0]["content"] = [{"type": "fill", "attrs": {"fid": "fill001", "label": "Placeholder"}}]
        repo.save_tab_doc(self.conn, self.slug, "shots", doc=doc, title="Shots", actor_id=self.actor, session_id="fixture", base_revision=None)
        self.conn.commit()
        base = self.tab()
        url = f"/projects/api/{self.slug}/fields"
        self.assertEqual(self.request(url, method="POST", data={"tab_key": "shots", "field_id": "fill001", "expected_value": "", "value": "other"})[0], 200)
        code, conflict, _ = self.edit(base, [("b_shot1", {"0": "mine", "1": "blocked"})])
        self.assertEqual(code, 409, conflict)
        self.assertEqual(node_text(self.tab()["doc"]["content"][0]["content"][1]["content"][1]), "original")
        self.assertEqual(self.request(url, method="POST", data={"tab_key": "shots", "field_id": "fill001", "expected_value": "", "value": "stale"})[0], 409)

    def test_schema_dependency_and_type_validation_precede_write(self):
        base = self.tab()
        table = base["doc"]["content"][0]
        url = f"/projects/api/{self.slug}/tabs/shots/table-schema"
        body = {"table_id": "b_table", "base_revision": base["revision"], "expected_column_ids": column_ids(table), "column_id": column_ids(table)[1]}
        code, result, _ = self.request(url, method="POST", data={**body, "operation": "change_type", "column_type": "number"})
        self.assertEqual(code, 400, result)
        code, edited, _ = self.edit(base, [("b_shot1", {"1": "=A1"})])
        self.assertEqual(code, 200, edited)
        code, blocked, _ = self.request(url, method="POST", data={**body, "base_revision": edited["revision"], "operation": "reorder", "order": list(reversed(column_ids(table)))})
        self.assertEqual(code, 400, blocked)
        self.assertTrue(blocked["dependencies"])
        self.assertEqual(column_ids(self.tab()["doc"]["content"][0]), column_ids(table))

    def test_xlsx_structure_requires_explicit_new_workbook_delivery(self):
        book = openpyxl.Workbook()
        book.active.append(["Name", "Value"])
        book.active.append(["A", "B"])
        output = io.BytesIO()
        book.save(output)
        raw = output.getvalue()
        headers = {"X-File-Name": "test.xlsx", "Content-Type": file_handoff.MIME["xlsx"]}
        base = f"/projects/api/{self.slug}"
        code, preview, _ = self.request(base + "/tabs/import-preview", method="POST", raw=raw, headers=headers)
        self.assertEqual(code, 200, preview)
        code, imported, _ = self.request(base + "/tabs/import", method="POST", raw=raw, headers=headers)
        self.assertEqual(code, 200, imported)
        tab = imported["tab"]
        payload = {"base_revision": tab["revision"], "table_id": tab["doc"]["content"][0]["attrs"]["bid"], "new_rows": [["C", "D"]]}
        url = base + f"/tabs/{tab['key']}"
        self.assertEqual(self.request(url + "/table-edit", method="POST", data=payload)[0], 400)
        code, result, _ = self.request(url + "/table-edit", method="POST", data={**payload, "delivery_mode": "new_workbook"})
        self.assertEqual(code, 200, result)
        self.assertEqual(self.request(url + "/export?format=xlsx")[0], 400)
        code, fresh, _ = self.request(url + "/export?format=xlsx&mode=new_workbook")
        self.assertEqual(code, 200)
        self.assertEqual(openpyxl.load_workbook(io.BytesIO(fresh)).active["A3"].value, "C")
        self.assertEqual(self.request(url + "/export?mode=source")[1], raw)
        self.assertFalse(self.tab(tab["key"])["source_meta"]["original_patch_supported"])

    def test_reimport_requires_review_and_orphans_ambiguous_row_annotations(self):
        base = f"/projects/api/{self.slug}"
        headers = {"X-File-Name": "rows.csv", "Content-Type": "text/csv"}
        raw = b"ID,Text\nA,duplicate\nA,duplicate\nB,unique\n"
        code, created, _ = self.request(base + "/tabs/import", method="POST", raw=raw, headers=headers)
        self.assertEqual(code, 200, created)
        tab = created["tab"]
        rows = tab["doc"]["content"][0]["content"]
        duplicate_ids = {rows[i]["attrs"]["bid"] for i in (1, 2)}
        unique_id = rows[3]["attrs"]["bid"]
        for rid in duplicate_ids:
            code, annotation, _ = self.request(base + "/annotations", method="POST", data={"tab_key": tab["key"], "row": rid, "body": "Keep this row identity"})
            self.assertEqual(code, 200, annotation)
        incoming = b"ID,Text\nA,external-one\nB,unique\nA,external-two\n"
        query = f"?tab_key={tab['key']}&base_revision={tab['revision']}"
        code, preview, _ = self.request(base + "/tabs/import-preview" + query, method="POST", raw=incoming, headers=headers)
        self.assertEqual(code, 200, preview)
        self.assertEqual(preview["replacement"]["summary"], {"added": 2, "removed": 2, "changed": 0, "unchanged": 1})
        self.assertTrue(preview["replacement"]["ambiguities"])
        self.assertEqual(self.request(base + "/tabs/import" + query, method="POST", raw=incoming, headers=headers)[0], 400)
        code, replaced, _ = self.request(base + "/tabs/import" + query + "&confirm_replace=true", method="POST", raw=incoming, headers=headers)
        self.assertEqual(code, 200, replaced)
        current_rows = replaced["tab"]["doc"]["content"][0]["content"]
        self.assertTrue(duplicate_ids.isdisjoint({row["attrs"]["bid"] for row in current_rows}))
        self.assertEqual(current_rows[2]["attrs"]["bid"], unique_id)
        annotations = repo.list_annotations(self.conn, self.slug)
        old_comments = [ann for ann in annotations if ann.get("anchor", {}).get("rid") in duplicate_ids]
        self.assertEqual(len(old_comments), 2)
        self.assertTrue(all(ann.get("orphaned_at") for ann in old_comments))


if __name__ == "__main__":
    unittest.main()
